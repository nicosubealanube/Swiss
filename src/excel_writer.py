"""Escritor de registros en Excel (SharePoint vía Microsoft Graph API y archivo .xlsx local)."""

import json
from pathlib import Path
from typing import List, Optional, Union
import urllib.request
from xml.sax.saxutils import escape
import zipfile
import xml.etree.ElementTree as ET

import config
from src.models import EXCEL_COLUMNS, ExcelRowRecord


class LocalExcelWriter:
    """
    Gestiona la lectura del último 'Número de ingreso' y la adición de nuevas filas
    en un archivo Excel (.xlsx) local o en una carpeta sincronizada de SharePoint/OneDrive.
    Soporta openpyxl si está disponible y cuenta con un motor nativo .xlsx sin dependencias externas.
    """

    def __init__(self, excel_path: Union[str, Path] = config.LOCAL_EXCEL_PATH) -> None:
        self.excel_path = Path(excel_path)

    def get_next_ingreso_number(self, default_base: int = config.INITIAL_INGRESO_NUMBER) -> int:
        """Obtiene el último 'Número de ingreso' del archivo Excel y le suma 1."""
        rows = self.read_all_rows()
        if len(rows) <= 1:
            return default_base + 1
        max_num = default_base
        for row in rows[1:]:
            if not row:
                continue
            try:
                val = int(float(str(row[0]).strip()))
                if val > max_num:
                    max_num = val
            except ValueError:
                pass
        return max_num + 1

    def read_all_rows(self) -> List[List[Union[str, int, float]]]:
        """Lee todas las filas del archivo .xlsx (incluyendo la cabecera)."""
        if not self.excel_path.exists():
            return []

        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append(["" if c is None else c for c in row])
            return rows
        except ImportError:
            return self._read_xlsx_native()

    def append_record(self, record: ExcelRowRecord) -> Path:
        """Agrega un registro al archivo .xlsx (creándolo con cabeceras si no existe)."""
        self.excel_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            import openpyxl  # type: ignore

            if self.excel_path.exists():
                wb = openpyxl.load_workbook(self.excel_path)
                ws = wb.active
            else:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Excepciones"
                ws.append(EXCEL_COLUMNS)
            ws.append(record.to_list())
            wb.save(self.excel_path)
            return self.excel_path
        except ImportError:
            all_rows = self._read_xlsx_native()
            if not all_rows:
                all_rows = [list(EXCEL_COLUMNS)]
            all_rows.append(record.to_list())
            self._write_xlsx_native(all_rows)
            return self.excel_path

    def _read_xlsx_native(self) -> List[List[Union[str, int, float]]]:
        """Lector nativo de .xlsx basado en librería estándar (zipfile + xml)."""
        if not self.excel_path.exists():
            return []
        ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(self.excel_path, "r") as zf:
            shared_strings: List[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                ss_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for si in ss_root.findall("main:si", ns):
                    texts = [t.text or "" for t in si.findall(".//main:t", ns)]
                    shared_strings.append("".join(texts))

            sheet_root = ET.fromstring(zf.read("xl/worksheets/sheet1.xml"))
            rows: List[List[Union[str, int, float]]] = []
            for row_el in sheet_root.findall(".//main:sheetData/main:row", ns):
                row_vals: List[Union[str, int, float]] = []
                for c_el in row_el.findall("main:c", ns):
                    c_type = c_el.attrib.get("t", "")
                    v_el = c_el.find("main:v", ns)
                    is_el = c_el.find("main:is/main:t", ns)
                    if c_type == "inlineStr" and is_el is not None:
                        row_vals.append(is_el.text or "")
                    elif c_type == "s" and v_el is not None and v_el.text is not None:
                        idx = int(v_el.text)
                        row_vals.append(shared_strings[idx] if idx < len(shared_strings) else "")
                    elif v_el is not None and v_el.text is not None:
                        raw_v = v_el.text
                        try:
                            row_vals.append(int(raw_v) if "." not in raw_v else float(raw_v))
                        except ValueError:
                            row_vals.append(raw_v)
                    else:
                        row_vals.append("")
                rows.append(row_vals)
            return rows

    def _write_xlsx_native(self, rows: List[List[Union[str, int, float, None]]]) -> None:
        """Generador nativo de archivo OpenXML .xlsx válido sin dependencias externas."""

        def col_letter(n: int) -> str:
            res = ""
            while n > 0:
                n, rem = divmod(n - 1, 26)
                res = chr(65 + rem) + res
            return res

        sheet_rows_xml = []
        for r_idx, row in enumerate(rows, start=1):
            cells_xml = []
            for c_idx, val in enumerate(row, start=1):
                ref = f"{col_letter(c_idx)}{r_idx}"
                style_attr = ' s="1"' if r_idx == 1 else ""
                if val is None or val == "":
                    cells_xml.append(f'<c r="{ref}" t="inlineStr"{style_attr}><is><t></t></is></c>')
                elif isinstance(val, (int, float)) and not isinstance(val, bool):
                    cells_xml.append(f'<c r="{ref}"{style_attr}><v>{val}</v></c>')
                else:
                    escaped = escape(str(val))
                    cells_xml.append(
                        f'<c r="{ref}" t="inlineStr"{style_attr}><is><t>{escaped}</t></is></c>'
                    )
            sheet_rows_xml.append(f'<row r="{r_idx}">{"".join(cells_xml)}</row>')

        sheet_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetData>{"".join(sheet_rows_xml)}</sheetData>'
            "</worksheet>"
        )

        content_types_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            "</Types>"
        )

        rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        )

        workbook_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Excepciones" sheetId="1" r:id="rId1"/></sheets>'
            "</workbook>"
        )

        workbook_rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            "</Relationships>"
        )

        styles_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="2">'
            '<font><sz val="10"/><name val="Calibri"/></font>'
            '<font><b/><sz val="10"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font>'
            "</fonts>"
            '<fills count="3">'
            '<fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill>'
            '<fill><patternFill patternType="solid"><fgColor rgb="FF1F4E78"/></patternFill></fill>'
            "</fills>"
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="2">'
            '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
            '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/>'
            "</cellXfs>"
            "</styleSheet>"
        )

        with zipfile.ZipFile(self.excel_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types_xml)
            zf.writestr("_rels/.rels", rels_xml)
            zf.writestr("xl/workbook.xml", workbook_xml)
            zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels_xml)
            zf.writestr("xl/styles.xml", styles_xml)
            zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)


class SharePointExcelWriter:
    """
    Cliente para insertar filas directamente en una Tabla de un Excel alojado en SharePoint
    utilizando Microsoft Graph API (permite escritura concurrente sin bloquear el archivo).
    """

    def __init__(
        self,
        access_token: str,
        site_id: str = config.SHAREPOINT_SITE_ID,
        drive_id: str = config.SHAREPOINT_DRIVE_ID,
        item_id: str = config.SHAREPOINT_ITEM_ID,
        table_name: str = config.SHAREPOINT_TABLE_NAME,
    ) -> None:
        self.access_token = access_token
        self.site_id = site_id
        self.drive_id = drive_id
        self.item_id = item_id
        self.table_name = table_name

    def _base_table_url(self) -> str:
        return (
            f"https://graph.microsoft.com/v1.0/sites/{self.site_id}"
            f"/drives/{self.drive_id}/items/{self.item_id}"
            f"/workbook/tables/{self.table_name}"
        )

    def get_next_ingreso_number(self, default_base: int = config.INITIAL_INGRESO_NUMBER) -> int:
        """Consulta las filas de la tabla en SharePoint para obtener el último Número de ingreso + 1."""
        url = f"{self._base_table_url()}/rows?$select=values"
        req = urllib.request.Request(
            url,
            headers={"Authorization": f"Bearer {self.access_token}"},
            method="GET",
        )
        with urllib.request.urlopen(req) as resp:
            payload = json.loads(resp.read().decode("utf-8"))

        max_num = default_base
        for row_obj in payload.get("value", []):
            values = row_obj.get("values", [[]])[0]
            if values:
                try:
                    num = int(float(str(values[0])))
                    if num > max_num:
                        max_num = num
                except ValueError:
                    pass
        return max_num + 1

    def append_record(self, record: ExcelRowRecord) -> dict:
        """Agrega una fila al final de la tabla del Excel en SharePoint."""
        url = f"{self._base_table_url()}/rows/add"
        body = json.dumps({"values": [record.to_list()]}).encode("utf-8")
        req = urllib.request.Request(
            url,
            data=body,
            headers={
                "Authorization": f"Bearer {self.access_token}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(req) as resp:
            return json.loads(resp.read().decode("utf-8"))
